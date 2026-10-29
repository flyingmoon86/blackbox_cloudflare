from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
REF = Path(r'C:\Users\gbg61\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-design-report\assets\reference.docx')
OUT = ROOT / 'artifacts/technical-archive/黑匣子 Cloudflare 网站技术归档.docx'

OUT.parent.mkdir(parents=True, exist_ok=True)

def plain(s):
    s = re.sub(r'\[([^]]+)\]\(([^)]+)\)', lambda m: f'{m.group(1)}（{m.group(2)}）', s)
    return s.replace('**', '').replace('`', '')

doc = Document(REF)
body = doc._element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)
for name in ('Title', 'Subtitle', 'Normal', 'Heading 1', 'Heading 2'):
    if name in doc.styles:
        st = doc.styles[name]
        st.font.name = 'Microsoft YaHei'
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].font.color.rgb = RGBColor(38, 32, 30)
doc.styles['Heading 1'].font.color.rgb = RGBColor(98, 31, 48)
doc.styles['Heading 2'].font.color.rgb = RGBColor(98, 31, 48)
doc.styles['Normal'].font.size = Pt(9.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.32
for sec in doc.sections:
    sec.top_margin = Inches(.78)
    sec.bottom_margin = Inches(.68)
    for node in sec.header._element.iter(qn('w:t')):
        if node.text:
            node.text = node.text.replace('Report title', '黑匣子技术归档').replace('Date', '2026 09 26')
    for p in sec.header.paragraphs:
        for run in p.runs:
            run.text = run.text.replace('Report title', '黑匣子技术归档').replace('Date', '2026 09 26')
    for table in sec.header.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.text = run.text.replace('Report title', '黑匣子技术归档').replace('Date', '2026 09 26')

def add_table(rows):
    rows = [[plain(x.strip()) for x in line.strip('|').split('|')] for line in rows]
    if len(rows) > 1 and all(re.fullmatch(r'[\s:-]+', x) for x in rows[1]):
        rows.pop(1)
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    if len(rows[0]) == 3:
        widths = [Inches(1.70), Inches(2.50), Inches(2.30)] if rows[0][0] == '请求' else [Inches(1.30), Inches(2.70), Inches(2.50)]
    else:
        widths = [Inches(6.5/len(rows[0]))] * len(rows[0])
    for ri, values in enumerate(rows):
        for ci, value in enumerate(values):
            cell = table.cell(ri, ci)
            cell.width = widths[ci]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.18
            p.paragraph_format.keep_with_next = ri < len(rows) - 1
            run = p.add_run(value)
            run.font.name = 'Microsoft YaHei'
            run.font.size = Pt(8.7 if ri else 9)
            if ri == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '6A2031' if ri == 0 else ('F7F1EB' if ri % 2 else 'FFFFFF'))
            cell._tc.get_or_add_tcPr().append(shade)
        tr_pr = table.rows[ri]._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        tr_pr.append(no_split)
        if ri == 0:
            tr_pr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph('')

in_code = False
lines = (ROOT / 'README.md').read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    i += 1
    if not line.strip():
        continue
    if line.startswith('|'):
        group = [line]
        while i < len(lines) and lines[i].startswith('|'):
            group.append(lines[i])
            i += 1
        add_table(group)
        continue
    if line.startswith('```'):
        in_code = not in_code
        continue
    if in_code:
        p = doc.add_paragraph(style='Normal')
        p.paragraph_format.left_indent = Inches(.18)
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(8)
        continue
    if line.startswith('# '):
        doc.add_paragraph(plain(line[2:]), style='Title')
    elif line.startswith('## '):
        p = doc.add_paragraph(plain(line[3:]), style='Heading 1')
        if line.startswith('## 二 ') or line.startswith('## 写法参考'):
            p.paragraph_format.page_break_before = True
    elif line.startswith('### '):
        doc.add_paragraph(plain(line[4:]), style='Heading 2')
    elif line.startswith('> '):
        p = doc.add_paragraph(plain(line[2:]), style='Normal')
        p.paragraph_format.left_indent = Inches(.2)
        p.paragraph_format.right_indent = Inches(.2)
    else:
        doc.add_paragraph(plain(line), style='Normal')
doc.save(OUT)
print(OUT)
